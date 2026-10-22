"""
One-time migration: export your edits from a shuffle-sheet workbook (v2.5) into a JSON the web app imports.

Reads the input cells only (not calculated values):
  typed prices (Reagent Prices C:E), vendor-buy curation (column K, stored as differences from game data),
  Settings, Intermediates "Allow crafting?", Recipes "Use chains? = No", the DE Table, and the Imports rows.
Items are keyed by item ID (the workbook keys by name; the Item ID column supplies the mapping).

Run:
  python3 tools/export_workbook_edits.py data/private/Vendor_DE_Master_Sheet_v2.5.xlsx
  -> data/private/edits-v2.5.json  (git-ignored; load it with the app's Import button)
"""
import json, os, sys
from openpyxl import load_workbook
from openpyxl.worksheet.formula import ArrayFormula

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
src = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, 'data', 'private', 'Vendor_DE_Master_Sheet_v2.5.xlsx')
dst = sys.argv[2] if len(sys.argv) > 2 else os.path.join(ROOT, 'data', 'private', 'edits-v2.5.json')
game = json.load(open(os.path.join(ROOT, 'app', 'public', 'game-data.json')))
lit = lambda v: v.text if isinstance(v, ArrayFormula) else v
wb = load_workbook(src)
hdr = lambda ws: {str(c.value).strip(): c.column - 1 for c in ws[1] if c.value is not None}

P = wb['Reagent Prices']; h = hdr(P); id_by_name = {}; typed = {}; vb = {}
for r in P.iter_rows(min_row=2, values_only=True):
    name, iid = r[h['Item']], r[h['Item ID']]
    if not name or iid is None: continue
    id_by_name[name] = iid
    g, s, c = (lit(r[h[k]]) for k in ('Gold', 'Silver', 'Copper'))
    if any(isinstance(x, (int, float)) for x in (g, s, c)):
        typed[str(iid)] = round((g or 0) * 10000 + (s or 0) * 100 + (c or 0))
    k = lit(r[h['Vendor buy price (ref)']]); k = None if k in ('', None) else k
    gd = game['items'].get(str(iid), {}).get('buy')
    used_as = str(r[h['Used as']] or '')
    if k != gd and ('Reagent' in used_as or 'DE material' in used_as):  # product-only rows are blanked by the build, not by you
        vb[str(iid)] = k  # null = "no vendor sells this"

S = wb['Settings']; settings = {}
for r in S.iter_rows(min_row=2, max_col=2, values_only=True):
    if r[0] and not str(r[0]).startswith('Intermediates tab ID') and not str(r[0]).startswith('Prices are per'):
        settings[r[0]] = lit(r[1])

X = wb['Intermediates']; h = hdr(X); allow = {}
for r in X.iter_rows(min_row=2, values_only=True):
    if r[h['Item']]: allow[str(id_by_name.get(r[h['Item']], r[h['Item']]))] = lit(r[h['Allow crafting?']])

W = wb['Recipes']; h = hdr(W)
no_chain = sorted(r[h['Recipe ID']] for r in W.iter_rows(min_row=2, values_only=True) if lit(r[h['Use chains?']]) == 'No')

T = wb['DE Table']; h = hdr(T); bands = []
for r in T.iter_rows(min_row=2, values_only=True):
    name = r[0]
    if not name or not isinstance(name, str) or name.startswith('Source:'): continue
    res = []
    for k in (1, 2, 3):
        m = lit(r[h[f'Result {k}']])
        if m: res.append({'item': id_by_name.get(m), 'name': m, 'chance': lit(r[h[f'Chance {k}']]) or 0,
                          'min': lit(r[h[f'Min qty {k}']]), 'max': lit(r[h[f'Max qty {k}']])})
    bands.append({'band': name, 'quality': r[h['Quality']], 'type': r[h['Item type']], 'ilvlMin': r[h['Item level min']],
                  'ilvlMax': r[h['Item level max']], 'results': res, 'note': r[h['Notes']] or ''})

I = wb['Imports']; imports = []
for r in I.iter_rows(min_row=2, max_col=5, values_only=True):
    if any(v not in (None, '') for v in r): imports.append({'item': r[0], 'name': r[1], 'p1': r[2], 'p2': r[3], 'p3': r[4]})

out = {'format': 'shuffle-edits', 'version': 1, 'exportedFrom': os.path.basename(src), 'typedPrices': typed,
       'vendorBuyOverrides': vb, 'settings': settings, 'allowCrafting': allow, 'useChainsNo': no_chain,
       'deTable': bands, 'imports': imports}
json.dump(out, open(dst, 'w'), indent=1)
print(f'{len(typed)} typed prices, {len(vb)} vendor-buy overrides, {len(settings)} settings, {len(allow)} intermediates, '
      f'{len(no_chain)} Use chains = No, {len(bands)} DE bands, {len(imports)} import rows -> {os.path.relpath(dst, ROOT)}')
