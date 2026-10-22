"""
Export the shuffle sheet's calculated results as the parity reference the web app is tested against.

Needs a workbook WITH calculated values: download it from Google Sheets (File > Download > .xlsx), or recalculate
a copy with LibreOffice. A file straight from the build script has formulas only and no values.

Run:
  python3 tools/export_parity.py PATH_TO_CALCULATED.xlsx
  -> data/private/parity-v2.5.json  (git-ignored: built from your own price scan)

Per recipe (by Recipe ID): reagent cost (buy all and cheapest path), vendor, DE and AH values, profits, best exit,
crafted inputs. Per intermediate (by item ID): buy price, craft cost, best cost, source. Money is in copper.
"""
import json, os, sys
from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
src = sys.argv[1]; dst = sys.argv[2] if len(sys.argv) > 2 else os.path.join(ROOT, 'data', 'private', 'parity-v2.5.json')
wb = load_workbook(src, data_only=True)
hdr = lambda ws: {str(c.value).strip(): c.column - 1 for c in ws[1] if c.value is not None}
num = lambda v: v if isinstance(v, (int, float)) else None
txt = lambda v: v if isinstance(v, str) and v != '' else None

W = wb['Recipes']; h = hdr(W)
if all(r[h['Reagents priced']] is None for r in W.iter_rows(min_row=2, max_row=50, values_only=True)):  # always "N of M" once calculated
    sys.exit('No calculated values in this workbook. Download it from Google Sheets as .xlsx or recalculate it first.')
cols = {'buyAll': 'Your reagent cost (buy all)', 'cheapest': 'Reagent cost (cheapest path)', 'vendorValue': 'Vendor value per craft',
        'deValue': 'Expected DE value (after AH cut)', 'ahPrice': 'AH price (each)', 'ahValue': 'AH value per craft (after AH cut)',
        'profitVendor': 'Profit vendoring', 'profitDE': 'Profit disenchanting', 'profitAH': 'Profit selling on AH',
        'profitBest': 'Profit at best exit'}
recipes = {}
for r in W.iter_rows(min_row=2, values_only=True):
    rid = r[h['Recipe ID']]
    if rid is None: continue
    d = {k: num(r[h[c]]) for k, c in cols.items()}
    d |= {'bestExit': txt(r[h['Best exit']]), 'craftedInputs': txt(r[h['Crafted inputs']]), 'deBand': txt(r[h['DE band']]), 'name': r[h['Recipe']]}
    recipes[str(rid)] = d
P = wb['Reagent Prices']; hp = hdr(P); ids = {r[hp['Item']]: r[hp['Item ID']] for r in P.iter_rows(min_row=2, values_only=True) if r[0]}
X = wb['Intermediates']; hx = hdr(X); inter = {}
for r in X.iter_rows(min_row=2, values_only=True):
    if r[0]: inter[str(ids.get(r[0], r[0]))] = {'name': r[0], 'buy': num(r[hx['Buy price (each)']]), 'craft': num(r[hx['Craft cost (each)']]),
                                               'best': num(r[hx['Best cost (each)']]), 'source': txt(r[hx['Source']])}
json.dump({'from': os.path.basename(src), 'recipes': recipes, 'intermediates': inter}, open(dst, 'w'), indent=1)
print(f'{len(recipes)} recipes, {len(inter)} intermediates -> {os.path.relpath(dst, ROOT)}')
